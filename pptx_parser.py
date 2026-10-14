"""
PPTX parser – extracts slide text and images from a PowerPoint file.

Each slide is represented as a ``Slide`` dataclass carrying the title,
body (as a markdown-like bullet list), full text content, and paths to
any extracted images.
"""

from __future__ import annotations

import base64
import logging
import math
from decimal import Decimal
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree

from pptx import Presentation
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

from config import ensure_temp_dir

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------


@dataclass
class Slide:
    """Represents a single parsed slide.

    Attributes:
        index: 1-based slide number.
        title: Extracted title text (empty string if none).
        body: Body text formatted as a markdown bullet list.
        image_paths: Absolute paths to extracted image files.
        text_content: All textual content on the slide (suitable for TTS / subtitles).
        notes: Speaker notes, also included in text_content but never in body.
    """

    index: int
    title: str = ""
    body: str = ""
    image_paths: list[str] = field(default_factory=list)
    text_content: str = ""
    notes: str = ""
    source_path: str = ""
    hidden: bool = False
    content_units: list[dict] = field(default_factory=list)
    page_width: int = 0
    page_height: int = 0


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _safe_shape_type(shape: BaseShape) -> int | None:
    """Safely get shape type, returning None for unrecognized shapes."""
    try:
        return shape.shape_type
    except NotImplementedError:
        return None
    except Exception:
        return None


def _is_title_shape(shape: BaseShape) -> bool:
    """Return True if *shape* is a PPTX title placeholder."""
    try:
        if shape.is_placeholder:
            ph = shape.placeholder_format  # type: ignore[attr-defined]
            # TITLE, CENTER_TITLE, VERTICAL_TITLE.
            return ph.type in (1, 3, 5)
    except Exception:
        pass
    return False


def _is_body_shape(shape: BaseShape) -> bool:
    """Return True if *shape* is a PPTX body / content placeholder."""
    try:
        if shape.is_placeholder:
            ph = shape.placeholder_format  # type: ignore[attr-defined]
            # 2 = BODY, 7 = OBJECT.
            return ph.type in (2, 7)
    except Exception:
        pass
    return False


def _extract_image(shape: Picture, out_dir: Path, slide_idx: int, img_idx: int) -> str:
    """Save the image from *shape* to *out_dir* and return its absolute path."""
    image = shape.image
    ext = image.content_type.split("/")[-1]
    if ext == "jpeg":
        ext = "jpg"
    filename = f"slide{slide_idx:02d}_img{img_idx:02d}.{ext}"
    filepath = out_dir / filename
    with open(filepath, "wb") as f:
        f.write(image.blob)
    return str(filepath.resolve())


def _shape_to_bullet_text(shape: BaseShape) -> str:
    """Convert a shape's text frames into a markdown-like bullet list."""
    if not shape.has_text_frame:
        return ""

    tf = shape.text_frame
    lines: list[str] = []
    for para in tf.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        level = para.level if para.level is not None else 0
        indent = "  " * level
        lines.append(f"{indent}- {text}")
    return "\n".join(lines)


_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_DIAGRAM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_IDENTITY = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)
_EMU_PER_PIXEL = 914400 / 96  # Display size uses the standard 96-DPI convention.


def _compose_transform(outer: tuple, inner: tuple) -> tuple:
    """Compose affine transforms, applying *inner* before *outer*."""
    a, b, c, d, e, f = outer
    g, h, i, j, k, l = inner
    return (
        a * g + c * h, b * g + d * h,
        a * i + c * j, b * i + d * j,
        a * k + c * l + e, b * k + d * l + f,
    )


def _orientation_transform(cx, cy, rotation=0, flip_h=False, flip_v=False):
    """Rotate and flip around a center in the containing coordinate system."""
    angle = math.radians(rotation)
    cosine, sine = math.cos(angle), math.sin(angle)
    fx, fy = (-1 if flip_h else 1), (-1 if flip_v else 1)
    a, b, c, d = cosine * fx, sine * fx, -sine * fy, cosine * fy
    return a, b, c, d, cx - a * cx - c * cy, cy - b * cx - d * cy


def _group_transform(shape: GroupShape) -> tuple:
    """Map a group's child coordinates to its parent's coordinates."""
    xfrm = shape._element.grpSpPr.xfrm
    if xfrm is None:
        return _IDENTITY
    off, ext, ch_off, ch_ext = xfrm.off, xfrm.ext, xfrm.chOff, xfrm.chExt
    left, top = (off.x, off.y) if off is not None else (0, 0)
    child_left, child_top = (ch_off.x, ch_off.y) if ch_off is not None else (0, 0)
    width, height = (ext.cx, ext.cy) if ext is not None else (0, 0)
    sx = width / ch_ext.cx if ch_ext is not None and ch_ext.cx else 1.0
    sy = height / ch_ext.cy if ch_ext is not None and ch_ext.cy else 1.0
    scale = (sx, 0, 0, sy, left - child_left * sx, top - child_top * sy)
    orientation = _orientation_transform(
        left + width / 2, top + height / 2,
        float(xfrm.get("rot", "0")) / 60000,
        xfrm.get("flipH") in ("1", "true"),
        xfrm.get("flipV") in ("1", "true"),
    )
    return _compose_transform(orientation, scale)


def _iter_shapes(shapes, transform=_IDENTITY):
    """Yield leaf shapes with their complete group-to-slide transform."""
    for shape in shapes:
        if isinstance(shape, GroupShape) or _safe_shape_type(shape) == 6:
            try:
                child_transform = _compose_transform(transform, _group_transform(shape))
            except Exception as exc:
                logger.warning("Could not read group coordinates: %s", exc)
                child_transform = transform
            try:
                yield from _iter_shapes(shape.shapes, child_transform)
            except Exception as exc:
                logger.warning("Could not read group contents: %s", exc)
        else:
            yield shape, transform


def _shape_geometry(shape: BaseShape, transform: tuple) -> tuple:
    """Return slide-space left/top, displayed edge lengths and displayed area."""
    left, top = float(shape.left or 0), float(shape.top or 0)
    width, height = float(shape.width or 0), float(shape.height or 0)
    orientation = _orientation_transform(
        left + width / 2, top + height / 2, float(shape.rotation or 0)
    )
    a, b, c, d, e, f = _compose_transform(transform, orientation)
    corners = [
        (a * x + c * y + e, b * x + d * y + f)
        for x, y in ((left, top), (left + width, top),
                     (left, top + height), (left + width, top + height))
    ]
    return (
        min(x for x, _ in corners), min(y for _, y in corners),
        math.hypot(a * width, b * width), math.hypot(c * height, d * height),
        abs(a * d - b * c) * width * height,
    )


def _table_to_text(shape: BaseShape) -> str:
    """Keep cells on the same row together, including multiline cell content."""
    lines = []
    for row in shape.table.rows:
        cells = []
        for cell in row.cells:
            try:
                text = "" if cell.is_spanned else cell.text
                cells.append(" / ".join(text.splitlines()).strip())
            except Exception as exc:
                logger.warning("Could not read table cell: %s", exc)
                cells.append("")
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _xml_text_lines(element) -> list[str]:
    """Read DrawingML paragraphs without splitting individual formatted runs."""
    paragraphs = list(element.iter(f"{{{_DRAWING_NS}}}p"))
    if paragraphs:
        return [text for paragraph in paragraphs if (text := "".join(
            node.text or "" for node in paragraph.iter(f"{{{_DRAWING_NS}}}t")
        ).strip())]
    return [node.text.strip() for node in element.iter(f"{{{_DRAWING_NS}}}t")
            if node.text and node.text.strip()]


def _chart_to_text(shape: BaseShape) -> str:
    """Extract series/category/value data; use XML caches for unsupported plots."""
    chart = shape.chart
    lines = []
    if chart.has_title and chart.chart_title.has_text_frame:
        title = chart.chart_title.text_frame.text.strip()
        if title:
            lines.append(title)
    for plot in chart.plots:
        try:
            categories = [" / ".join(str(label) for label in labels if label is not None)
                          for labels in plot.categories.flattened_labels]
        except (AttributeError, ValueError, NotImplementedError):
            categories = []
        for series in plot.series:
            try:
                name = series.name or "Series"
                values = series.values
                # Scatter/bubble charts store their X axis in c:xVal, not c:cat.
                x_values = _chart_cache_values(series._element, "xVal")
                labels = categories or x_values
                points = [f"{labels[i] if i < len(labels) else i + 1}: "
                          f"{format(Decimal(str(value)), 'f')}" if isinstance(value, (int, float)) else
                          f"{labels[i] if i < len(labels) else i + 1}: "
                          f"{value if value is not None else '—'}"
                          for i, value in enumerate(values)]
                lines.append(f"{name}: {'; '.join(points)}" if points else name)
            except Exception as exc:
                logger.warning("Could not read chart series; using XML cache: %s", exc)
                lines.extend(_chart_xml_lines(series._element))
    if not lines:
        lines = _chart_xml_lines(chart._element)
    return "\n".join(lines)


def _chart_cache_values(element, tag: str) -> list[str]:
    """Read indexed cached points, preserving gaps in category/value positions."""
    container = element.find(f"{{{_CHART_NS}}}{tag}")
    if container is None:
        return []
    points = {}
    for point in container.iter(f"{{{_CHART_NS}}}pt"):
        value = point.find(f"{{{_CHART_NS}}}v")
        if value is not None and value.text is not None:
            points[int(point.get("idx", len(points)))] = value.text
    return [points.get(i, "") for i in range(max(points, default=-1) + 1)]


def _chart_xml_lines(element) -> list[str]:
    lines = []
    for series in element.iter(f"{{{_CHART_NS}}}ser"):
        names = _chart_cache_values(series, "tx")
        if not names:
            names = [node.text for node in series.findall(
                f"{{{_CHART_NS}}}tx/{{{_CHART_NS}}}v") if node.text]
        labels = _chart_cache_values(series, "cat") or _chart_cache_values(series, "xVal")
        values = _chart_cache_values(series, "val") or _chart_cache_values(series, "yVal")
        points = [f"{labels[i] if i < len(labels) else i + 1}: {value or '—'}"
                  for i, value in enumerate(values)]
        if names or points:
            lines.append(f"{' / '.join(names) or 'Series'}: {'; '.join(points)}")
    return lines


def _smartart_to_text(shape: BaseShape) -> str:
    """Follow SmartArt's r:dm relationship to its actual diagram data part."""
    lines = []
    for rel_ids in shape._element.iter(f"{{{_DIAGRAM_NS}}}relIds"):
        rid = rel_ids.get(f"{{{_REL_NS}}}dm")
        if not rid:
            continue
        try:
            part = shape.part.related_part(rid)
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            root = etree.fromstring(part.blob, parser=parser)
            part_lines = _xml_text_lines(root)
            # Some producers store rendered text only in the diagram drawing part.
            if not part_lines:
                for relation in part.rels.values():
                    if relation.is_external or not relation.reltype.endswith("/diagramDrawing"):
                        continue
                    root = etree.fromstring(relation.target_part.blob, parser=parser)
                    part_lines.extend(_xml_text_lines(root))
            lines.extend(part_lines)
        except Exception as exc:
            logger.warning("Could not read SmartArt data relationship %s: %s", rid, exc)
    return "\n".join(lines or _xml_text_lines(shape._element))


def _extract_notes(slide) -> str:
    """Read existing speaker notes without creating a new notes slide."""
    try:
        if not slide.has_notes_slide:
            return ""
        text_frame = slide.notes_slide.notes_text_frame
        return text_frame.text.strip() if text_frame is not None else ""
    except Exception as exc:
        logger.warning("Could not read speaker notes: %s", exc)
        return ""


def _usable_image_area(shape: Picture, geometry: tuple) -> float:
    """Reject small source/display edges; return area used to rank the survivors."""
    display_width, display_height, area = geometry[2:]
    if min(display_width, display_height) / _EMU_PER_PIXEL < 80 - 1e-6:
        return 0
    try:
        if min(shape.image.size) < 80:
            return 0
    except Exception as exc:
        # Vector images and unusual image formats may have no pixel dimensions.
        logger.debug("Image pixel dimensions unavailable, using display size: %s", exc)
    return area


def image_to_base64_data_uri(image_path: str) -> str:
    """Read an image file and return a base64 data-URI string.

    Args:
        image_path: Absolute path to an image file.

    Returns:
        A string like ``"data:image/png;base64,iVBOR..."`` suitable for
        use in SiliconFlow API payloads.
    """
    path = Path(image_path)
    ext = path.suffix.lstrip(".").lower()
    if ext == "jpg":
        ext = "jpeg"
    mime = f"image/{ext}"
    with open(path, "rb") as f:
        b64 = base64.b64encode(f.read()).decode("ascii")
    return f"data:{mime};base64,{b64}"


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def parse_pptx(pptx_path: str | Path, *, include_hidden: bool = False) -> list[Slide]:
    """Parse a PPTX file and return a list of :class:`Slide` objects.

    Args:
        pptx_path: Path to a ``.pptx`` file.

    Returns:
        List of slides extracted from the presentation, in order.
    """
    pptx_path = Path(pptx_path).resolve()
    if not pptx_path.exists():
        raise FileNotFoundError(f"PPTX file not found: {pptx_path}")
    if pptx_path.suffix.lower() != ".pptx":
        raise ValueError(f"Expected a .pptx file, got: {pptx_path.suffix}")

    temp_dir = ensure_temp_dir()
    images_dir = temp_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(pptx_path))
    slides: list[Slide] = []

    logger.info("Parsing %s (%d slides)", pptx_path.name, len(prs.slides))

    for slide_num, slide in enumerate(prs.slides, start=1):
        hidden = slide._element.get('show', '1') in {'0', 'false'}
        if hidden and not include_hidden:
            continue
        content_units = []
        text_blocks: list[tuple[str, str]] = []
        title_candidates: list[tuple[int, float, float, int]] = []
        image_candidates: list[tuple[float, Picture]] = []
        image_paths: list[str] = []
        table_count = 0

        for shape, transform in _iter_shapes(slide.shapes):
            kind = 'text'
            try:
                geometry = _shape_geometry(shape, transform)
            except Exception as exc:
                logger.debug("Slide %d: shape position unavailable: %s", slide_num, exc)
                geometry = (math.inf, math.inf, 0, 0, 0)
            try:
                if _safe_shape_type(shape) == 13 or isinstance(shape, Picture):
                    area = _usable_image_area(shape, geometry)
                    if area:
                        image_candidates.append((area, shape))
                    continue

                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if not text:
                        continue
                    bullets = _shape_to_bullet_text(shape)
                    is_title = _is_title_shape(shape)
                    left, top = geometry[:2]
                    if is_title or top < prs.slide_height * 0.15:
                        title_candidates.append((
                            0 if is_title else 1, top, left, len(text_blocks)
                        ))
                elif shape.has_table:
                    kind = 'table'
                    table_count += 1
                    text = _table_to_text(shape)
                    bullets = "\n".join(f"- {line}" for line in text.splitlines())
                elif shape.has_chart:
                    kind = 'chart'
                    try:
                        text = _chart_to_text(shape)
                    except Exception as exc:
                        logger.warning("Slide %d: using chart XML fallback: %s", slide_num, exc)
                        text = "\n".join(_chart_xml_lines(shape.chart._element))
                    bullets = "\n".join(f"- {line}" for line in text.splitlines())
                else:
                    text = _smartart_to_text(shape)
                    bullets = "\n".join(f"- {line}" for line in text.splitlines())
                if text:
                    text_blocks.append((text, bullets))
                    for line_index, line in enumerate(text.splitlines()):
                        if line.strip():
                            role = 'title' if _is_title_shape(shape) else ('footer' if geometry[1] > prs.slide_height * .9 else 'body')
                            content_units.append(dict(id=f'{slide_num}:{shape.shape_id}:{line_index}',
                                text=line.strip(), source='title' if role == 'title' else 'body',
                                role=role, kind=kind, left=geometry[0], top=geometry[1],
                                width=geometry[2], height=geometry[3]))
            except Exception as exc:
                logger.warning("Slide %d: failed to read shape: %s", slide_num, exc)

        title_index = min(title_candidates)[3] if title_candidates else None
        title = text_blocks[title_index][0] if title_index is not None else ""
        other_blocks = [block for i, block in enumerate(text_blocks) if i != title_index]
        body = "\n".join(bullets for _, bullets in other_blocks if bullets)
        notes = _extract_notes(slide)
        content_units.extend(dict(id=f'{slide_num}:notes:{i}', text=line.strip(),
                                  source='notes', role='notes', kind='text')
                             for i, line in enumerate(notes.splitlines()) if line.strip())
        all_text_parts = ([title] if title else []) + [text for text, _ in other_blocks]
        if notes:
            all_text_parts.append(notes)
        text_content = "\n".join(all_text_parts)

        # The first images are the most prominent ones, for callers selecting [:N].
        for img_idx, (_, shape) in enumerate(
            sorted(image_candidates, key=lambda item: item[0], reverse=True), start=1
        ):
            try:
                image_paths.append(_extract_image(shape, images_dir, slide_num, img_idx))
            except Exception as exc:
                logger.warning("Slide %d: failed to extract image: %s", slide_num, exc)

        slide_obj = Slide(
            index=slide_num,
            title=title,
            body=body,
            image_paths=image_paths,
            text_content=text_content,
            notes=notes,
            source_path=str(pptx_path), hidden=hidden, content_units=content_units,
            page_width=int(prs.slide_width), page_height=int(prs.slide_height),
        )
        slides.append(slide_obj)
        logger.info(
            "Slide %d: title=%r, body_lines=%d, tables=%d, images=%d, notes_len=%d, text_len=%d",
            slide_num,
            title[:60] if title else "(none)",
            sum(bool(line.strip()) for line in body.splitlines()),
            table_count,
            len(image_paths),
            len(notes),
            len(text_content),
        )

    return slides
