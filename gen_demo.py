"""Generate a demo PPTX for testing pptx-to-video."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Slide 1 - Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(2))
tf = txBox.text_frame
tf.text = "智能产品演示"
tf.paragraphs[0].font.size = Pt(44)

txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(3))
tf2 = txBox2.text_frame
tf2.text = "2025 年度新品发布\nAI 驱动的下一代解决方案"
tf2.paragraphs[0].font.size = Pt(28)

# Slide 2 - Features
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
txBox3 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(1.5))
tf3 = txBox3.text_frame
tf3.text = "核心功能"
tf3.paragraphs[0].font.size = Pt(36)

txBox4 = slide2.shapes.add_textbox(Inches(1), Inches(3), Inches(11), Inches(4))
tf4 = txBox4.text_frame
for line in [
    "• 智能识别与理解",
    "• 毫秒级响应速度",
    "• 多语言自然交互",
    "• 安全可靠的架构",
]:
    p = tf4.add_paragraph()
    p.text = line
    p.font.size = Pt(24)

prs.save(str(Path(__file__).resolve().parent / "demo.pptx"))
print("demo.pptx created")
