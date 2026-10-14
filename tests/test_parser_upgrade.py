"""Offline regression coverage for content missed by the original PPTX parser."""

from __future__ import annotations

import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.util import Inches

import pptx_parser


FIXTURES = Path(__file__).parent / "fixtures"


def _textbox(shapes, text, left, top, width=4, height=0.4):
    shape = shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.text = text
    return shape


def _picture(shapes, size, color, left, top, width, height):
    stream = io.BytesIO()
    Image.new("RGB", size, color).save(stream, format="PNG")
    stream.seek(0)
    return shapes.add_picture(
        stream, Inches(left), Inches(top), Inches(width), Inches(height)
    )


def _diagram_frame(slide, rid, shape_id):
    frame = parse_xml(
        f'''<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
             xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
             xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
             xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
          <p:nvGraphicFramePr><p:cNvPr id="{shape_id}" name="Diagram {shape_id}"/>
            <p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>
          <p:xfrm><a:off x="457200" y="1828800"/>
            <a:ext cx="4572000" cy="1828800"/></p:xfrm>
          <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">
            <dgm:relIds r:dm="{rid}"/>
          </a:graphicData></a:graphic>
        </p:graphicFrame>'''
    )
    slide.shapes._spTree.insert_element_before(frame, "p:extLst")


def build_complex_pptx(path):
    """Build real groups, charts, notes and diagram relationships without Office."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(6)

    # Ordinary text boxes, including two candidates inside the top 15 percent.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide.shapes, "Quarterly review", 0.5, 0.1)
    _textbox(slide.shapes, "A subtitle that must survive", 0.5, 0.65)
    box = _textbox(slide.shapes, "Ordinary body", 0.5, 1.6, height=1.2)
    paragraph = box.text_frame.add_paragraph()
    paragraph.text = "Nested bullet"
    paragraph.level = 1
    slide.notes_slide.notes_text_frame.text = "Speaker only\nExplain the assumptions."

    # Actual table and category chart, neither exposes a normal text frame.
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Metrics"
    table = slide.shapes.add_table(
        3, 2, Inches(0.5), Inches(1.5), Inches(3), Inches(2)
    ).table
    for row, values in enumerate((("Metric", "Value"), ("Revenue", "42"), ("Users", "900"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("North", (10, 20))
    data.add_series("South", (15, 25))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(1.5),
        Inches(5), Inches(3.5), data
    )

    # The child title has top=4 inches, but is shown near the slide's top.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    _textbox(group.shapes, "Grouped heading", 2, 4)
    _textbox(group.shapes, "Grouped body", 2, 6)
    nested = group.shapes.add_group_shape()
    _textbox(nested.shapes, "Nested group detail", 2, 8)
    group.left, group.top = Inches(0.5), Inches(0.1)
    group.width, group.height = Inches(8), Inches(4.4)

    # A child near local y=0 must not become a slide title when its group is low.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    _textbox(group.shapes, "Low group body", 0, 0)
    group.left, group.top = Inches(0.5), Inches(4)
    _textbox(slide.shapes, "Real top heading", 0.5, 0.2)

    # Insertion order deliberately puts a smaller image first.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide.shapes, "Images", 0.5, 0.1)
    _picture(slide.shapes, (800, 600), "blue", 0.5, 2, 2, 2)
    _picture(slide.shapes, (1000, 800), "red", 3, 2, 4, 3)
    _picture(slide.shapes, (40, 200), "green", 7, 2, 2, 2)
    _picture(slide.shapes, (200, 40), "cyan", 7, 3, 2, 2)
    _picture(slide.shapes, (500, 500), "yellow", 7, 4, 0.5, 2)
    _picture(slide.shapes, (500, 500), "orange", 7, 4, 2, 0.5)
    _picture(slide.shapes, (80, 80), "white", 7, 4, 80 / 96, 80 / 96)
    group = slide.shapes.add_group_shape()
    _picture(group.shapes, (300, 300), "black", 0, 0, 2, 2)
    group.left, group.top = Inches(8), Inches(4)
    group.width, group.height = Inches(0.5), Inches(0.5)

    # A SmartArt GraphicFrame points to a diagram data part containing a:t.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _textbox(slide.shapes, "Workflow", 0.5, 0.1)
    part = Part(
        PackURI("/ppt/diagrams/data1.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
        slide.part.package, (FIXTURES / "smartart_data.xml").read_bytes()
    )
    rid = slide.part.relate_to(part, RT.DIAGRAM_DATA)
    _diagram_frame(slide, rid, 3)
    # A broken diagram must not suppress the following usable text.
    _diagram_frame(slide, "rIdMissingDiagram", 4)
    _textbox(slide.shapes, "After broken diagram", 0.5, 4.5)

    # Centered titles are not in the top-15-percent band.
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Centered title"
    slide.placeholders[1].text = "Centered subtitle"
    slide.shapes.title.top = Inches(2)

    prs.save(path)


class ParserUpgradeTests(unittest.TestCase):
    def test_chart_extraction_keeps_more_than_six_significant_digits(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ['Revenue', 'Ratio']
        data.add_series('Exact values', (1234567, 0.123456789))
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), data)
        text = pptx_parser._chart_to_text(frame)
        self.assertIn('1234567', text)
        self.assertIn('0.123456789', text)
        self.assertNotIn('1.23457e', text)

    @classmethod
    def setUpClass(cls):
        cls.scratch = tempfile.TemporaryDirectory(prefix="parser-upgrade-")
        cls.root = Path(cls.scratch.name)
        cls.pptx_path = cls.root / "complex.pptx"
        build_complex_pptx(cls.pptx_path)
        with patch.object(pptx_parser, "ensure_temp_dir", return_value=cls.root):
            cls.slides = pptx_parser.parse_pptx(cls.pptx_path)

    @classmethod
    def tearDownClass(cls):
        cls.scratch.cleanup()

    def test_top_band_title_keeps_other_candidates_in_body(self):
        slide = self.slides[0]
        self.assertEqual(slide.title, "Quarterly review")
        self.assertIn("- A subtitle that must survive", slide.body)
        self.assertNotIn("Quarterly review", slide.body)

    def test_ordinary_text_and_paragraph_levels_are_in_body(self):
        self.assertIn("- Ordinary body\n  - Nested bullet", self.slides[0].body)

    def test_table_rows_are_in_body_and_full_text(self):
        slide = self.slides[1]
        for value in ("Metric", "Value", "Revenue", "42", "Users", "900"):
            self.assertIn(value, slide.body)
            self.assertIn(value, slide.text_content)
        self.assertTrue(any("Revenue" in line and "42" in line for line in slide.body.splitlines()))

    def test_chart_categories_series_and_values_are_in_body(self):
        slide = self.slides[1]
        for value in ("North", "South", "Q1", "Q2", "10", "20", "15", "25"):
            self.assertIn(value, slide.body)
            self.assertIn(value, slide.text_content)

    def test_group_child_coordinates_are_transformed_and_recursed(self):
        slide = self.slides[2]
        self.assertEqual(slide.title, "Grouped heading")
        self.assertIn("Grouped body", slide.body)
        self.assertIn("Nested group detail", slide.body)

    def test_low_group_does_not_supply_a_false_title(self):
        slide = self.slides[3]
        self.assertEqual(slide.title, "Real top heading")
        self.assertIn("Low group body", slide.body)

    def test_images_filter_each_small_edge_and_sort_by_display_area(self):
        paths = self.slides[4].image_paths
        self.assertEqual(len(paths), 3)
        sizes = []
        for path in paths:
            self.assertTrue(Path(path).is_absolute())
            with Image.open(path) as picture:
                sizes.append(picture.size)
        self.assertEqual(sizes, [(1000, 800), (800, 600), (80, 80)])

    def test_smartart_follows_data_relationship_and_joins_text_runs(self):
        slide = self.slides[5]
        for value in ("Research", "Build prototype", "Launch"):
            self.assertIn(value, slide.body)
            self.assertIn(value, slide.text_content)
        self.assertEqual(slide.text_content.count("Build prototype"), 1)

    def test_broken_graphic_does_not_hide_following_body(self):
        self.assertIn("After broken diagram", self.slides[5].body)

    def test_center_title_placeholder_is_recognized(self):
        self.assertEqual(self.slides[6].title, "Centered title")
        self.assertIn("Centered subtitle", self.slides[6].body)

    def test_speaker_notes_are_separate_and_in_full_text_only(self):
        slide = self.slides[0]
        self.assertEqual(getattr(slide, "notes", None), "Speaker only\nExplain the assumptions.")
        self.assertIn("Speaker only\nExplain the assumptions.", slide.text_content)
        self.assertNotIn("Speaker only", slide.body)
        self.assertEqual(getattr(self.slides[1], "notes", None), "")

    def test_parsing_does_not_create_missing_notes_slides(self):
        presentation = Presentation(self.pptx_path)
        without_notes = [slide for slide in presentation.slides if not slide.has_notes_slide]
        with patch.object(pptx_parser, "Presentation", return_value=presentation), patch.object(
            pptx_parser, "ensure_temp_dir", return_value=self.root
        ):
            pptx_parser.parse_pptx(self.pptx_path)
        self.assertTrue(without_notes)
        self.assertTrue(all(not slide.has_notes_slide for slide in without_notes))

    def test_log_includes_per_slide_content_counts(self):
        with patch.object(pptx_parser, "ensure_temp_dir", return_value=self.root), self.assertLogs(
            "pptx_parser", level="INFO"
        ) as logs:
            pptx_parser.parse_pptx(self.pptx_path)
        slide_logs = [line for line in logs.output if "title=" in line]
        self.assertEqual(len(slide_logs), 7)
        for line in slide_logs:
            for field in ("body_lines=", "tables=", "images=", "notes_len="):
                self.assertIn(field, line)
        self.assertIn("tables=1", slide_logs[1])
        self.assertIn("images=3", slide_logs[4])

    def test_slide_positional_compatibility_and_default_notes(self):
        slide = pptx_parser.Slide(1, "Title", "Body", ["image.png"], "Full text")
        self.assertEqual(slide.text_content, "Full text")
        self.assertEqual(getattr(slide, "notes", None), "")

    def test_existing_image_data_uri_helper(self):
        value = pptx_parser.image_to_base64_data_uri(self.slides[4].image_paths[0])
        self.assertTrue(value.startswith("data:image/png;base64,"))


if __name__ == "__main__":
    unittest.main()
