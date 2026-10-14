import unittest
from types import SimpleNamespace
from pptx_parser import Slide
from faithful_plan import build_faithful_plans, audit_faithful_plans
from video_composer import build_scene_timeline


class ScriptFlowTests(unittest.TestCase):
    def test_audio_cache_reuses_exact_script_but_invalidates_changed_text(self):
        import asyncio
        import tempfile
        from pathlib import Path
        from unittest.mock import AsyncMock, patch
        import tts_audio
        with tempfile.TemporaryDirectory() as folder:
            async def synth(text, output_path=None, **kwargs):
                Path(output_path).write_bytes(b'valid-test-audio')
                return Path(output_path)
            scene = dict(index=1, text='原稿。', preserve_text=True, cache_dir=folder)
            with patch.object(tts_audio, '_edge_tts', new=AsyncMock(side_effect=synth)) as generate, patch.object(tts_audio, '_audio_duration', return_value=3):
                asyncio.run(tts_audio.generate_scene_narrations([scene], Path(folder)/'first'))
                asyncio.run(tts_audio.generate_scene_narrations([scene], Path(folder)/'second'))
                self.assertEqual(generate.await_count, 1)
                scene['text'] = '修改稿。'
                asyncio.run(tts_audio.generate_scene_narrations([scene], Path(folder)/'third'))
                self.assertEqual(generate.await_count, 2)

    def test_parser_keeps_page_source_and_hidden_metadata(self):
        import tempfile
        from pathlib import Path
        from pptx import Presentation
        from pptx_parser import parse_pptx
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'deck.pptx'
            prs = Presentation()
            first = prs.slides.add_slide(prs.slide_layouts[1])
            first.shapes.title.text = '标题'
            first.placeholders[1].text = '收入 10%，仅限试点'
            hidden = prs.slides.add_slide(prs.slide_layouts[1])
            hidden._element.set('show', '0')
            prs.save(path)
            slides = parse_pptx(path)
            self.assertEqual(len(slides), 1)
            self.assertEqual(slides[0].source_path, str(path.resolve()))
            self.assertTrue(slides[0].content_units)
            self.assertEqual(len(parse_pptx(path, include_hidden=True)), 2)

    def test_final_edited_script_survives_refresh_and_audit(self):
        from gui import _refresh_plan_narration
        slides = [Slide(1, '原标题', '- 原内容')]
        plan = build_faithful_plans(slides)[0]
        plan.narration_state = 'edited'
        plan.narration_text = plan.subtitle_text = '这是我修改后的讲解。'
        plan.narration_by_slide = {1: plan.narration_text}
        _refresh_plan_narration(slides, [plan], 'explain')
        self.assertEqual(plan.narration_text, '这是我修改后的讲解。')
        audit = audit_faithful_plans(slides, [plan])
        self.assertEqual(audit['pages'][0]['narration_verification'], 'user_edited_unverified')
        self.assertEqual(build_scene_timeline(slides, [plan])[0]['text'], plan.narration_text)

    def test_explicit_silence_never_falls_back_to_slide(self):
        slides = [Slide(1, '不应朗读')]
        plan = SimpleNamespace(slide_index=1, slide_indices=[1], narration_text='',
                               narration_by_slide={1: ''}, narration_state='edited',
                               narration_muted_by_slide={1: True})
        self.assertEqual(build_scene_timeline(slides, [plan])[0]['text'], '')

    def test_merged_scene_preserves_each_final_segment_exactly(self):
        slides = [Slide(1, '标题一'), Slide(2, '标题二')]
        plan = SimpleNamespace(slide_index=1, slide_indices=[1, 2], narration_text='甲乙',
                               narration_by_slide={1: '甲', 2: '乙'}, narration_state='generated')
        scenes = build_scene_timeline(slides, [plan])
        self.assertEqual([s['text'] for s in scenes], ['甲', '乙'])
        self.assertTrue(all(s['preserve_text'] for s in scenes))


if __name__ == '__main__':
    unittest.main()
