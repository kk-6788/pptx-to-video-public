import json
from unittest.mock import patch

import unittest
import tempfile
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement


def deck(tmp_path):
    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if i == 1:
            slide._element.set('show', '0')
        if i == 2:
            slide._element.append(OxmlElement('p:timing'))
    path = tmp_path / 'source.pptx'
    prs.save(path)
    return path


def export(source, directory, slide_indices):
    for i in slide_indices:
        Image.new('RGB', (640, 360), (i * 50, 10, 20)).save(directory / f'slide-{i}.png')


def test_hidden_numbering_and_animation_warning(tmp_path):
    from pptx_page_renderer import render_pptx_pages
    with patch('pptx_page_renderer._export_powerpoint', side_effect=export):
        pages = render_pptx_pages(deck(tmp_path), tmp_path / 'pages')
    assert list(pages) == [1, 3]
    assert pages[3]['slide_index'] == 3
    assert any('动画' in warning for warning in pages[3]['warnings'])
    assert (pages[1]['width'], pages[1]['height']) == (640, 360)
    assert len(pages[1]['source_sha256']) == 64


def test_cache_reuses_only_verified_images_and_source(tmp_path):
    from pptx_page_renderer import render_pptx_pages
    source, out = deck(tmp_path), tmp_path / 'pages'
    with patch('pptx_page_renderer._export_powerpoint', side_effect=export) as renderer:
        first = render_pptx_pages(source, out, include_hidden=True)
        assert list(first) == [1, 2, 3]
        assert render_pptx_pages(source, out, include_hidden=True) == first
        assert renderer.call_count == 1
        Image.new('RGB', (640, 360), 'white').save(first[1]['path'])
        render_pptx_pages(source, out, include_hidden=True)
        assert renderer.call_count == 2
        prs = Presentation(source)
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(source)
        assert len(render_pptx_pages(source, out, include_hidden=True)) == 4
        assert renderer.call_count == 3


def test_partial_export_never_publishes_cache(tmp_path):
    from pptx_page_renderer import render_pptx_pages
    def partial(source, directory, indices):
        export(source, directory, indices[:1])
    with patch('pptx_page_renderer._export_powerpoint', side_effect=partial):
        with unittest.TestCase().assertRaises((ValueError, FileNotFoundError)):
            render_pptx_pages(deck(tmp_path), tmp_path / 'pages')
    assert not list((tmp_path / 'pages').rglob('manifest.json'))


def test_renderer_failure_is_clear_not_card_fallback(tmp_path):
    from pptx_page_renderer import render_pptx_pages, PageRenderError
    with patch('pptx_page_renderer._export_powerpoint', side_effect=PageRenderError('PowerPoint unavailable')):
        with unittest.TestCase().assertRaisesRegex(PageRenderError, 'PowerPoint'):
            render_pptx_pages(deck(tmp_path), tmp_path / 'pages')


class PageRendererTests(unittest.TestCase):
    pass


for _name, _fn in list(globals().items()):
    if _name.startswith('test_'):
        def _run(self, fn=_fn):
            with tempfile.TemporaryDirectory() as tmp:
                fn(Path(tmp))
        setattr(PageRendererTests, _name, _run)
del _name, _fn
