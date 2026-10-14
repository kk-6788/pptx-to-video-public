import asyncio
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import main as cli


class EntrypointUpgradeTests(unittest.TestCase):
    def test_provider_override_applies_only_to_its_own_stage(self):
        for flag, expected in [('--image-provider', ('openai', 'kling')),
                               ('--video-provider', ('siliconflow', 'openai'))]:
            with self.subTest(flag=flag), patch.dict(os.environ, {
                'PROVIDER': 'agnes', 'IMAGE_PROVIDER': 'siliconflow', 'VIDEO_PROVIDER': 'kling',
            }), patch.object(cli, '_run_pipeline', new=AsyncMock(return_value=0)) as pipeline:
                self.assertEqual(cli.main(['-i', 'demo.pptx', '-o', 'out.mp4', flag, 'openai']), 0)
                options = pipeline.await_args.kwargs
                self.assertEqual((options['image_provider_name'], options['video_provider_name']), expected)

    def test_cli_accepts_format_narration_and_audio_options(self):
        args = cli._build_parser().parse_args([
            '-i', 'deck.pptx', '-o', 'out.mp4', '--ratio', '9:16',
            '--narration-style', 'explain', '--tts', '--bgm', 'music.mp3',
            '--no-bgm', '--srt-only',
        ])
        self.assertEqual(args.ratio, '9:16')
        self.assertEqual(args.narration_style, 'explain')
        self.assertTrue(args.add_audio)
        self.assertTrue(args.no_bgm)
        self.assertTrue(args.srt_only)
        self.assertEqual(args.bgm, Path('music.mp3'))

    def test_explicit_provider_wins_over_environment(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'deck.pptx'
            source.touch()
            run = AsyncMock(return_value=0)
            with patch.dict(os.environ, {
                'PROVIDER': 'siliconflow', 'IMAGE_PROVIDER': 'openai',
                'VIDEO_PROVIDER': 'kling',
            }), patch.object(cli, '_run_pipeline', run):
                self.assertEqual(cli.main([
                    '-i', str(source), '-o', str(source.with_suffix('.mp4')),
                    '--provider', 'comfyui',
                ]), 0)
            self.assertEqual(run.call_args.kwargs['provider_name'], 'comfyui')
            self.assertEqual(run.call_args.kwargs['image_provider_name'], 'comfyui')
            self.assertEqual(run.call_args.kwargs['video_provider_name'], 'comfyui')

    def test_srt_only_exports_local_planned_subtitles_without_network(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'deck.pptx'
            destination = Path(folder) / 'export' / 'out.mp4'
            presentation = Presentation()
            for title, body in [('增长趋势', '转化率提升35%'), ('优化流程', '服务成本降低20%')]:
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = title
                slide.placeholders[1].text = body
            presentation.save(source)
            forbidden = AssertionError('offline mode called a network/generation boundary')
            with patch('providers.get_providers', side_effect=forbidden), \
                    patch('video_generator.generate_clips_from_plans', side_effect=forbidden), \
                    patch('tts_audio.generate_narration', side_effect=forbidden), \
                    patch('video_composer.compose_video', side_effect=forbidden):
                code = cli.main([
                    '-i', str(source), '-o', str(destination), '--srt-only',
                    '--narration-style', 'explain', '--ratio', '9:16',
                ])
            self.assertEqual(code, 0)
            subtitles = destination.with_suffix('.srt')
            self.assertTrue(subtitles.exists())
            self.assertFalse(destination.exists())
            data = subtitles.read_bytes()
            self.assertTrue(data.startswith(b'\xef\xbb\xbf'))
            text = data.decode('utf-8-sig')
            self.assertIn('-->', text)
            self.assertIn('35%', text)
            self.assertIn('20%', text)

    def test_dry_run_stops_before_director_and_composition(self):
        from pptx_parser import Slide
        with patch('pptx_parser.parse_pptx', return_value=[Slide(1, '标题')]), \
                patch('prompt_engine.Director', side_effect=AssertionError('director invoked')), \
                patch('video_composer.compose_video', side_effect=AssertionError('composer invoked')):
            self.assertEqual(asyncio.run(cli._run_pipeline(
                Path('deck.pptx'), Path('out.mp4'), fps=24, add_subtitles=False,
                max_concurrent=1, provider_name='comfyui', image_model='', video_model='',
                style='auto', overall_theme='', dry_run=True, preview_only=False,
            )), 0)

    def test_composition_runs_outside_active_event_loop_and_receives_options(self):
        from pptx_parser import Slide
        slide = Slide(1, '标题', text_content='转化率提升35%')
        plan = SimpleNamespace(slide_index=1, director_vision='vision', image_prompt='image',
                               video_prompt='video')
        provider = SimpleNamespace(close=AsyncMock())
        received = {}

        def compose(*args, **kwargs):
            asyncio.run(asyncio.sleep(0))
            received.update(kwargs)

        with patch('pptx_parser.parse_pptx', return_value=[slide]), \
                patch('prompt_engine.Director') as director, \
                patch('providers.get_providers', return_value=(provider, provider)), \
                patch('video_generator.generate_clips_from_plans', new=AsyncMock(return_value=[])), \
                patch('video_generator.cleanup_temp_files', new=AsyncMock()), \
                patch('video_composer.compose_video', side_effect=compose):
            director.return_value.plan.return_value = [plan]
            self.assertEqual(asyncio.run(cli._run_pipeline(
                Path('demo.pptx'), Path('out.mp4'), fps=30, add_subtitles=True,
                max_concurrent=1, provider_name='comfyui', image_provider_name='comfyui',
                video_provider_name='comfyui', image_model='', video_model='', style='auto',
                overall_theme='', dry_run=False, preview_only=False, ratio='9:16',
                add_audio=True, add_bgm=False, bgm_path=Path('music.mp3'),
                narration_style='explain', visual_mode='creative',
            )), 0)
        self.assertEqual(received['plans'], [plan])
        self.assertTrue(received['add_audio'])
        self.assertFalse(received['add_bgm'])
        self.assertEqual(received['bgm_path'], Path('music.mp3'))
        self.assertEqual(received['narration_style'], 'explain')

    def test_gui_refreshes_grouped_narration_without_overwriting_visual_prompts(self):
        import gui
        from pptx_parser import Slide
        slides = [Slide(1, '增长', text_content='转化率提升35%'),
                  Slide(2, '成本', text_content='成本降低20%')]
        plan = SimpleNamespace(slide_index=1, slide_indices=[1, 2],
                               image_prompt='用户编辑的画面', video_prompt='用户编辑的运镜')
        with patch('tts_audio.build_slide_narration',
                   side_effect=lambda slide, narration_style: f'{narration_style}:{slide.text_content}',
                   create=True):
            gui._refresh_plan_narration(slides, [plan], 'explain')
        self.assertEqual(plan.narration_style, 'explain')
        self.assertEqual(plan.narration_by_slide[2], 'explain:成本降低20%')
        self.assertIn('35%', plan.narration_text)
        self.assertIn('20%', plan.narration_text)
        self.assertEqual(plan.image_prompt, '用户编辑的画面')
        self.assertEqual(plan.video_prompt, '用户编辑的运镜')


if __name__ == '__main__':
    unittest.main()
