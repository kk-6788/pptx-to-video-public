"""Reproducible real-file acceptance checks, separate from unit tests.

python tools/acceptance.py media|tts|full|gui|grouped
"""
from __future__ import annotations
import argparse
import asyncio
import importlib.util
import importlib.machinery
import json
import logging
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import config
import video_composer as vc
from pptx_parser import Slide, parse_pptx
from prompt_engine import Director, ScriptWriter
from video_generator import GeneratedClip
from PIL import Image

OUT = ROOT / 'artifacts' / 'acceptance'
OUT.mkdir(parents=True, exist_ok=True)


def probe(path):
    data = json.loads(subprocess.check_output([
        'ffprobe', '-v', 'error', '-show_streams', '-show_format', '-of', 'json', str(path)]))
    return {'duration': float(data['format']['duration']),
            'streams': [{k: s[k] for k in ('codec_type', 'codec_name', 'width', 'height',
                                          'sample_rate', 'channels', 'bit_rate') if k in s}
                        for s in data['streams']]}


def load_source(name, path):
    loader = importlib.machinery.SourceFileLoader(name, str(path))
    spec = importlib.util.spec_from_loader(name, loader)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    loader.exec_module(module)
    return module


def save_results(name, data):
    (OUT / f'{name}.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(data, ensure_ascii=False, indent=2), flush=True)


def media():
    config.set_ratio('16:9')
    slides = parse_pptx(ROOT / 'demo.pptx')
    old = load_source('old_parser_acceptance', ROOT / 'pptx_parser.py.bak')
    old_slides = old.parse_pptx(ROOT / 'demo.pptx')
    results = {'demo': [{'index': s.index, 'old_body_lines': len(o.body.splitlines()),
                         'body_lines': len(s.body.splitlines()), 'images': len(s.image_paths),
                         'notes_length': len(s.notes)} for o, s in zip(old_slides, slides)]}
    old_composer = load_source('old_composer_acceptance', ROOT / 'video_composer.py.bak')
    results['old_background_mean'] = float(old_composer._make_animated_bg_frame(0, 1280, 720, 5).mean())
    fallback_path = OUT / 'fallback_5s.mp4'
    clip = vc._make_fallback_clip(slides[-1], 5)
    clip.save_frame(str(OUT / 'fallback_landscape.png'), t=2.5)
    clip.write_videofile(str(fallback_path), fps=24, codec='libx264', audio=False, logger=None)
    clip.close()
    results['fallback'] = probe(fallback_path)
    samples = vc._detect_black_frames(fallback_path)
    results['fallback_brightness'] = samples
    assert samples and not vc._has_black_segment(samples) and min(b for _, b in samples) > 150

    black_path = OUT / 'black_3s.mp4'
    subprocess.run(['ffmpeg', '-y', '-v', 'error', '-f', 'lavfi', '-i',
                    'color=black:s=1280x720:r=24:d=3', '-c:v', 'libx264', str(black_path)], check=True)
    samples = vc._detect_black_frames(black_path)
    assert vc._has_black_segment(samples)
    results['black_samples'] = samples
    results['black_ranges'] = vc._black_segments(samples)
    vc._quality_check_video(black_path)
    assert not vc._has_black_segment(vc._detect_black_frames(black_path, sample_interval=.5), 4)

    # Exercise input black replacement and BGM without TTS at the same time.
    config.VIDEO_CODEC = 'libx264'
    replaced = vc.compose_video(slides, [GeneratedClip(slides[0].index, str(black_path), 3, 3)],
                                OUT / 'black_replaced_bgm.mp4', add_subtitles=True, fps=24)
    results['black_replaced'] = probe(replaced)
    assert any(s['codec_type'] == 'audio' for s in results['black_replaced']['streams'])
    assert not vc._has_black_segment(vc._detect_black_frames(replaced))
    # Check every transition frame, not only one sample per second.
    from moviepy import VideoFileClip
    with VideoFileClip(str(replaced)) as video:
        results['transition_brightness'] = [float(video.get_frame(2.85 + i / 24).mean()) for i in range(17)]
    assert min(results['transition_brightness']) > 150

    # Real PPT fixture supplements demo's absent table/SmartArt/image cases.
    fixture_module = load_source('parser_fixture_acceptance', ROOT / 'tests/test_parser_upgrade.py')
    fixture_module.build_complex_pptx(OUT / 'complex.pptx')
    config.TEMP_DIR = OUT / 'complex_media'
    complex_slides = parse_pptx(OUT / 'complex.pptx')
    assert 'Revenue' in complex_slides[1].body and 'Research' in complex_slides[5].body
    illustrated = Slide(1, '用数据看懂产品变化', '- 转化率提升了 35%\n- 服务成本下降了 20%\n- 下一步继续优化服务',
                        image_paths=complex_slides[4].image_paths)
    config.set_ratio('9:16')
    vertical = vc.compose_video([illustrated], [], OUT / 'out_9x16.mp4',
                                add_subtitles=True, fps=24)
    results['vertical'] = probe(vertical)
    assert results['vertical']['streams'][0]['width'] == 720
    assert results['vertical']['streams'][0]['height'] == 1280
    with VideoFileClip(str(vertical)) as video:
        video.save_frame(str(OUT / 'fallback_portrait.png'), t=1.8)
    results['bgm'] = {p.name: probe(p) for p in (ROOT / 'assets/bgm').glob('*.mp3')}
    save_results('media', results)


def tts():
    from pptx import Presentation
    from pptx.util import Inches
    from tts_audio import generate_scene_narrations
    prs = Presentation(str(ROOT / 'demo.pptx'))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.5), Inches(.1), Inches(8), Inches(.6))
    title.text = '经营结果验收样本'
    body = slide.shapes.add_textbox(Inches(.5), Inches(1.2), Inches(8), Inches(3))
    body.text = '转化率提升了35%，服务成本下降了20%。\n下一阶段继续优化服务，保持现有客户支持。'
    path = OUT / 'demo_plus_test_slide.pptx'
    prs.save(path)
    slides = parse_pptx(path)
    script = ScriptWriter(slides, narration_style='explain').write()
    by_slide = {index: text for scene in script.scenes for index, text in scene.narration_by_slide.items()}
    scenes = [{'index': s.index, 'text': by_slide[s.index], 'start_time': (s.index - 1) * 5,
               'duration': 5} for s in slides]
    results = asyncio.run(generate_scene_narrations(scenes, OUT / 'tts_verified'))
    assert len(results) == 3 and all(r and Path(r['path']).exists() and r['duration'] > 0 for r in results)
    for scene in scenes:
        assert not scene['text'].startswith(('首先', '其次', '最后'))
        assert 40 <= len(re.findall(r'[\u4e00-\u9fff]', scene['text'])) <= 80
    save_results('tts_verified', {'python': sys.version, 'source': str(path), 'scenes': scenes, 'audio': results})


def full():
    config.set_ratio('16:9')
    slides = parse_pptx(ROOT / 'demo.pptx')
    plans = Director(slides, narration_style='explain').plan()
    path = vc.compose_video(slides, [], OUT / 'out.mp4', plans=plans, add_subtitles=True,
                            add_audio=True, narration_style='explain')
    results = probe(path)
    assert any(s['codec_type'] == 'audio' for s in results['streams'])
    assert not vc._has_black_segment(vc._detect_black_frames(path))
    assert path.with_suffix('.srt').read_bytes().startswith(b'\xef\xbb\xbf')
    from moviepy import VideoFileClip
    with VideoFileClip(str(path)) as video:
        video.save_frame(str(OUT / 'full_subtitles.png'), t=video.duration * .65)
        # RMS confirms the encoded stream contains actual sound, not only silence.
        audio = video.audio.subclipped(1, min(4, video.duration)).to_soundarray(fps=8000)
        results['audio_rms'] = float((audio ** 2).mean() ** .5)
        assert results['audio_rms'] > .001
    save_results('full', results)


def grouped():
    """Encode one three-part generated video over three original slide clocks."""
    import numpy as np
    from types import SimpleNamespace
    from moviepy import VideoClip, VideoFileClip
    colors = [(240, 30, 30), (30, 240, 30), (30, 30, 240)]
    source = OUT / 'grouped_source.mp4'
    with VideoClip(lambda t: np.full((180, 320, 3), colors[min(2, int(t))], dtype=np.uint8),
                   duration=3) as clip:
        clip.write_videofile(str(source), fps=24, codec='libx264', audio=False, logger=None)
    slides = [Slide(i, f'分组场景第{i}页') for i in (1, 2, 3)]
    plan = SimpleNamespace(slide_index=1, slide_indices=[1, 2, 3], suggested_duration_sec=3)
    config.set_ratio('16:9')
    config.VIDEO_CODEC = 'libx264'
    output = vc.compose_video(slides, [GeneratedClip(1, str(source), 3, 3)],
                              OUT / 'grouped_verified.mp4', plans=[plan],
                              add_subtitles=True, add_bgm=False, fps=24)
    pixels = []
    with VideoFileClip(str(output)) as video:
        for second, channel in ((.65, 0), (1.65, 1), (2.65, 2)):
            pixel = video.get_frame(second)[360, 640]
            assert pixel[channel] > 200 and pixel[(channel + 1) % 3] < 60, (second, pixel)
            pixels.append({'time': second, 'rgb': pixel.tolist()})
    assert len(re.findall(r' --> ', output.with_suffix('.srt').read_text(encoding='utf-8-sig'))) == 3
    save_results('grouped', {'video': probe(output), 'samples': pixels, 'subtitle_cues': 3})


def gui():
    import gui as app_module
    from tkinter import Tk
    from PIL import ImageGrab
    root = Tk()
    app = app_module.PptxToVideoApp(root)
    errors = []
    root.report_callback_exception = lambda *args: errors.append(str(args))

    def check():
        try:
            root.update_idletasks()
            assert app.ratio_var.get() == '16:9'
            assert not app.narration_explain_var.get()
            assert app.bgm_var.get()
            assert app.bgm_volume_var.get() == 15
            app.ratio_var.set('9:16')
            app.narration_explain_var.set(True)
            app.bgm_volume_var.set(25)
            assert app._ratio_combo.cget('values') == ('16:9', '9:16')
            assert not errors, errors
            # Capture this app window only, even if another app has focus.
            import ctypes
            hwnd = ctypes.windll.user32.GetAncestor(root.winfo_id(), 2)
            ImageGrab.grab(window=hwnd).save(OUT / 'gui.png')
            save_results('gui', {'python': sys.version, 'startup': 'ok', 'ratio': app.ratio_var.get(),
                                 'explain': app.narration_explain_var.get(), 'bgm': app.bgm_var.get(),
                                 'volume': app.bgm_volume_var.get(), 'callback_errors': errors})
        except Exception as exc:
            errors.append(str(exc))
        finally:
            root.destroy()
    root.after(1800, check)
    root.mainloop()
    assert not errors, errors


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=['media', 'tts', 'full', 'gui', 'grouped'])
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format='%(levelname)s %(message)s', handlers=[
        logging.StreamHandler(), logging.FileHandler(OUT / f'{args.mode}.log', encoding='utf-8', mode='w')])
    globals()[args.mode]()
