"""Offline complex-PPT fixture, render and provenance/media verification."""
from pathlib import Path
import json
import subprocess
import sys
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import config
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont, ImageStat
from pptx_parser import parse_pptx
from faithful_plan import build_faithful_plans
from video_composer import compose_video, _detect_black_frames, _black_segments

OUT = ROOT/'artifacts/content-fidelity'
OUT.mkdir(parents=True, exist_ok=True)

def fixture():
    # This is a labeled test asset, not invented material in a user's slide.
    original = Image.new('RGB',(720,380),(229,236,234))
    draw = ImageDraw.Draw(original)
    font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',30)
    draw.rectangle((32,32,688,348), outline=(52,76,81), width=4)
    draw.text((60,85),'原图测试：产品结构示意',font=font,fill=(32,49,64))
    draw.text((60,175),'A → B → C',font=font,fill=(108,91,54))
    image = OUT/'fixture_source.png'
    original.save(image)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = '季度结果：保留数字与口径'
    table = slide.shapes.add_table(4,3,Inches(.5), Inches(1.5), Inches(9), Inches(3)).table
    for r, row in enumerate([['指标','数值','口径'],['营收','128.75万元','本季度'],
                              ['转化率','23.6%','网站访客'],['成本变化','-3.5%','同比']]):
        for c, value in enumerate(row):
            table.cell(r,c).text = value
    slide.notes_slide.notes_text_frame.text = '比较时保持同一统计口径。'
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = '处理顺序与原图'
    slide.placeholders[1].text = ('输入 → 校验 → 输出\n1. 接收订单\n2. 核对数量与金额\n3. 返回确认结果\n'
        '异常订单保留原因，禁止跳过校验。\n重复提交只处理一次，保留原始记录。')
    slide.shapes.add_picture(str(image), Inches(6), Inches(3), width=Inches(3))
    source = OUT/'complex_source.pptx'
    deck.save(source)
    slides = parse_pptx(source)
    config.set_ratio('16:9')
    config.VIDEO_CODEC = 'libx264'
    compose_video(slides, [], OUT/'complex.mp4', plans=build_faithful_plans(slides, 'explain'),
                  add_subtitles=True, add_bgm=True, fps=12, visual_mode='faithful')

def inspect(path):
    manifest = json.loads(path.with_suffix('.content.json').read_text(encoding='utf-8'))
    data = json.loads(subprocess.check_output(['ffprobe','-v','error','-show_streams','-show_format','-of','json',str(path)]))
    subprocess.run(['ffmpeg','-v','error','-i',str(path),'-f','null','-'],check=True,capture_output=True)
    samples = _detect_black_frames(path)
    assert samples and not _black_segments(samples)
    frames = []
    for i, scene in enumerate(manifest['scenes']):
        if i:
            assert abs(scene['start']-manifest['scenes'][i-1]['end']) < .001
        visual = scene['visuals']
        assert visual['readable_duration_satisfied']
        reconstructed = [''] * len(scene['source_facts'])
        for page in visual['pages']:
            for fragment in page['fragments']:
                reconstructed[fragment['fact_index']] += fragment['text']
        assert reconstructed == scene['source_facts']
        for label, at in [('middle',(scene['start']+scene['end'])/2),
                          ('begin',scene['start']+.05),('end',scene['end']-.25)]:
            output = OUT/f'{path.stem}_page{i+1}_{label}.png'
            subprocess.run(['ffmpeg','-v','error','-y','-ss',str(at),'-i',str(path),'-frames:v','1',str(output)],check=True)
            with Image.open(output) as image:
                brightness = ImageStat.Stat(image.convert('L')).mean[0]
            assert brightness > 150
            frames.append({'path':str(output),'time':at,'brightness':brightness})
    return {'path':str(path),'full_decode':'passed','duration':float(data['format']['duration']),
            'streams':[{k:s[k] for k in ['codec_type','codec_name','width','height'] if k in s} for s in data['streams']],
            'page_order':[s['slide_index'] for s in manifest['scenes']],
            'fact_count':sum(len(s['source_facts']) for s in manifest['scenes']),
            'all_source_facts_rendered':True,'minimum_brightness':min(b for _,b in samples),
            'black_segments':[], 'frames':frames}

if __name__=='__main__':
    if '--fixture' in sys.argv:
        fixture()
    else:
        results = [inspect(path) for path in [ROOT/'out.mp4',OUT/'out_9x16.mp4',OUT/'complex.mp4']]
        (OUT/'media-checks.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps(results,ensure_ascii=False,indent=2))
